#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
把 ZCode 开源公告做成高档学术风 PPT。

设计系统：
  纸底 #F7F5F1（暖白）· 深藏青 #14202E（标题）· 正文 #33414F · 弱化 #7C8896
  强调琥珀 #B87A2B · 次强调青 #1F6F7A · 分隔线 #D8D2C8
  标题衬线（华文中宋）· 正文无衬线（等线）· 数字英文 Georgia
  深色页底 #101A26（封面/结语）
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, ".."))
FIG = os.path.join(ROOT, "ppt_zcode", "figs")
OUTDIR = os.path.join(ROOT, "ppt_zcode")
os.makedirs(OUTDIR, exist_ok=True)

# ---------------- 设计令牌 ----------------
PAPER   = "F7F5F1"
INK     = "14202E"
BODY    = "33414F"
MUTED   = "7C8896"
AMBER   = "B87A2B"
TEAL    = "1F6F7A"
RULE    = "D8D2C8"
DARK    = "101A26"
WHITE   = "FFFFFF"
PAPER_D = "EDE9E2"

F_TITLE = "华文中宋"
F_SANS  = "等线"
F_SANSB = "等线"
F_SERIF = "Georgia"

SW = Inches(13.333)
SH = Inches(7.5)
M  = Inches(0.86)          # 页边距
CW = SW - 2 * M            # 内容宽


def set_font(run, size, color, bold=False, ea=F_SANS, lat=None, italic=False, spacing=None):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = RGBColor.from_string(color)
    f.name = lat or ea
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", ea)
    if spacing is not None:
        rPr.set("spc", str(int(spacing * 100)))


def tb(slide, l, t, w, h, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return box, tf


def para(tf, first=False):
    return tf.paragraphs[0] if first else tf.add_paragraph()


def line(slide, l, t, w, color=RULE, thick=Pt(0.75)):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, thick)
    s.fill.solid(); s.fill.fore_color.rgb = RGBColor.from_string(color)
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def blank(prs, bg=PAPER):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor.from_string(bg)
    return s


def page_head(slide, num, kicker, title, dark=False):
    """学术页眉：小号编号+栏目名 / 大标题 / 分隔线"""
    ink = WHITE if dark else INK
    mut = "9AA7B4" if dark else MUTED
    _, tf = tb(slide, M, Inches(0.52), CW, Inches(0.3))
    p = para(tf, True)
    r = p.add_run(); r.text = f"{num}   {kicker}"
    set_font(r, 11, mut, ea=F_SANS, lat=F_SANS, spacing=1.6)

    _, tf2 = tb(slide, M, Inches(0.92), CW, Inches(0.66))
    p2 = para(tf2, True)
    r2 = p2.add_run(); r2.text = title
    set_font(r2, 27, ink, bold=False, ea=F_TITLE, lat=F_TITLE)

    line(slide, M, Inches(1.72), CW, RULE if not dark else "2A3644")


def figure_caption(slide, l, t, w, text, dark=False):
    _, tf = tb(slide, l, t, w, Inches(0.4))
    p = para(tf, True)
    r = p.add_run(); r.text = text
    set_font(r, 10, "8B98A6" if dark else MUTED, ea=F_SANS, lat=F_SERIF, italic=True)


# ================================================================ 构建
def build():
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH

    # ---------------------------------------------------------- 1 封面
    s = blank(prs, DARK)
    cover = os.path.join(FIG, "01-cover.png")
    if os.path.exists(cover):
        s.shapes.add_picture(cover, 0, 0, width=SW, height=SH)
        # 左侧压深，保证文字可读
        ov = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(8.2), SH)
        ov.fill.solid(); ov.fill.fore_color.rgb = RGBColor.from_string("0B131C")
        ov.fill.transparency = 0.12
        ov.line.fill.background(); ov.shadow.inherit = False
        ov2 = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.4), 0, Inches(2.8), SH)
        ov2.fill.solid(); ov2.fill.fore_color.rgb = RGBColor.from_string("0B131C")
        ov2.fill.transparency = 0.55
        ov2.line.fill.background(); ov2.shadow.inherit = False

    _, tf = tb(s, M, Inches(0.78), Inches(7.0), Inches(0.4))
    p = para(tf, True)
    r = p.add_run(); r.text = "官方公告"
    set_font(r, 12, "C8A15E", ea=F_SANS, lat=F_SANS, spacing=3.0)

    _, tf = tb(s, M, Inches(1.35), Inches(7.4), Inches(1.5))
    p = para(tf, True)
    r = p.add_run(); r.text = "ZCode 开源"
    set_font(r, 60, WHITE, ea=F_TITLE, lat=F_TITLE, spacing=1.5)

    line(s, M, Inches(2.86), Inches(1.5), AMBER, Pt(2.2))

    _, tf = tb(s, M, Inches(3.14), Inches(6.9), Inches(1.2))
    p = para(tf, True)
    r = p.add_run(); r.text = "一次安全事件的整改、审计与开源"
    set_font(r, 21, "DCE3EA", ea=F_SANS, lat=F_SANS)
    p2 = para(tf)
    r2 = p2.add_run(); r2.text = "从社区反馈到两份第三方审计：公告说了什么，审计证实了什么"
    set_font(r2, 13, "8FA0AF", ea=F_SANS, lat=F_SANS)
    p2.space_before = Pt(10)

    _, tf = tb(s, M, Inches(6.42), Inches(7.4), Inches(0.5))
    p = para(tf, True)
    r = p.add_run(); r.text = "2026 年 9 月 21 日 08:45（北京）    依据：ZCode 官方公告原文"
    set_font(r, 11.5, "7E8C9B", ea=F_SANS, lat=F_SERIF, spacing=0.8)

    # ---------------------------------------------------------- 2 摘要
    s = blank(prs)
    page_head(s, "00", "摘要", "一页看懂")
    items = [
        ("整改完成并致歉", "针对社区反馈的 ZCode 产品安全问题，官方公告称已完成整改，并向所有用户道歉。"),
        ("代码已开源", "ZCode 已开源至 github.com/zai-org/ZCode，把代码交给社区监督，让产品开放、透明。"),
        ("两份第三方审计", "中国信息通信研究院与绿盟科技分别开展安全审计，结论指向 OSS 存储桶与 v3.14.0 客户端。"),
    ]
    y = Inches(2.08)
    for i, (t, d) in enumerate(items):
        _, tf = tb(s, M, y, Inches(0.6), Inches(0.5))
        p = para(tf, True)
        r = p.add_run(); r.text = f"0{i+1}"
        set_font(r, 26, RULE, ea=F_TITLE, lat=F_SERIF)

        _, tf = tb(s, M + Inches(0.86), y - Inches(0.02), CW - Inches(0.86), Inches(0.42))
        p = para(tf, True)
        r = p.add_run(); r.text = t
        set_font(r, 19, INK, bold=True, ea=F_SANS, lat=F_SANS)

        _, tf = tb(s, M + Inches(0.86), y + Inches(0.42), CW - Inches(1.4), Inches(0.6))
        p = para(tf, True)
        r = p.add_run(); r.text = d
        set_font(r, 13, BODY, ea=F_SANS, lat=F_SANS)
        p.line_spacing = 1.35

        y += Inches(1.42)
        if i < 2:
            line(s, M + Inches(0.86), y - Inches(0.24), CW - Inches(0.86))

    # ---------------------------------------------------------- 3 事件与回应
    s = blank(prs)
    page_head(s, "01", "事件", "社区反馈与官方回应")
    _, tf = tb(s, M, Inches(2.05), Inches(6.6), Inches(2.4))
    p = para(tf, True)
    r = p.add_run(); r.text = "起因"
    set_font(r, 12, AMBER, bold=True, ea=F_SANS, lat=F_SANS, spacing=2.0)
    p2 = para(tf)
    r2 = p2.add_run(); r2.text = "社区反馈的 ZCode 产品安全问题。"
    set_font(r2, 17, INK, ea=F_SANS, lat=F_SANS)
    p2.space_before = Pt(8)
    p3 = para(tf)
    r3 = p3.add_run(); r3.text = "公告原文致谢：非常感谢此前发现 ZCode 问题的社区开发者。"
    set_font(r3, 13, BODY, ea=F_SANS, lat=F_SANS)
    p3.space_before = Pt(26)

    p4 = para(tf)
    r4 = p4.add_run(); r4.text = "回应"
    set_font(r4, 12, AMBER, bold=True, ea=F_SANS, lat=F_SANS, spacing=2.0)
    p4.space_before = Pt(30)
    p5 = para(tf)
    r5 = p5.add_run(); r5.text = "已完成整改，并向所有用户道歉。"
    set_font(r5, 17, INK, ea=F_SANS, lat=F_SANS)
    p5.space_before = Pt(8)

    line(s, M + Inches(7.0), Inches(2.05), Inches(0.02), RULE)
    _, tf = tb(s, M + Inches(7.4), Inches(2.05), CW - Inches(7.4), Inches(3.0))
    p = para(tf, True)
    r = p.add_run(); r.text = "公告自报时间"
    set_font(r, 12, MUTED, ea=F_SANS, lat=F_SANS, spacing=2.0)
    p2 = para(tf)
    r2 = p2.add_run(); r2.text = "2026 年 9 月 21 日"
    set_font(r2, 24, INK, ea=F_TITLE, lat=F_SERIF)
    p2.space_before = Pt(6)
    p3 = para(tf)
    r3 = p3.add_run(); r3.text = "08:45 · 北京"
    set_font(r3, 15, AMBER, ea=F_SANS, lat=F_SERIF)
    p4 = para(tf)
    r4 = p4.add_run(); r4.text = "署名：原创 ZCode"
    set_font(r4, 12, MUTED, ea=F_SANS, lat=F_SANS)
    p4.space_before = Pt(16)

    # ---------------------------------------------------------- 4 整改清单
    s = blank(prs)
    page_head(s, "02", "整改", "已完成的三项整改")
    rows = [
        ("移除 Repo Wiki 功能", "Repo Wiki 入口及相应生成链路已移除。"),
        ("切断本地仓库快照链路", "已切断本地仓库快照生成与上传链路。"),
        ("客户端安全整改", "ZCode v3.14.0 客户端已完成安全整改。"),
    ]
    y = Inches(2.10)
    for i, (t, d) in enumerate(rows):
        box = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, M, y, CW, Inches(1.12))
        box.fill.solid(); box.fill.fore_color.rgb = RGBColor.from_string(PAPER_D)
        box.line.fill.background(); box.shadow.inherit = False
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, M, y, Pt(3.2), Inches(1.12))
        bar.fill.solid(); bar.fill.fore_color.rgb = RGBColor.from_string(AMBER)
        bar.line.fill.background(); bar.shadow.inherit = False

        _, tf = tb(s, M + Inches(0.34), y + Inches(0.20), CW - Inches(0.7), Inches(0.34))
        p = para(tf, True)
        r = p.add_run(); r.text = t
        set_font(r, 16.5, INK, bold=True, ea=F_SANS, lat=F_SANS)
        _, tf = tb(s, M + Inches(0.34), y + Inches(0.60), CW - Inches(0.7), Inches(0.34))
        p = para(tf, True)
        r = p.add_run(); r.text = d
        set_font(r, 12.5, BODY, ea=F_SANS, lat=F_SANS)
        y += Inches(1.36)

    # ---------------------------------------------------------- 5 开源
    s = blank(prs)
    page_head(s, "03", "开源", "把代码交给社区监督")
    fig = os.path.join(FIG, "02-opensource.png")
    if os.path.exists(fig):
        s.shapes.add_picture(fig, M + Inches(6.55), Inches(2.02), width=Inches(5.05))
        figure_caption(s, M + Inches(6.55), Inches(4.95), Inches(5.05),
                       "图 1 · 代码向社区展开：从单一仓库到可被持续检查的网络")

    _, tf = tb(s, M, Inches(2.12), Inches(6.0), Inches(3.2))
    p = para(tf, True)
    r = p.add_run(); r.text = "我们已将 ZCode 开源"
    set_font(r, 21, INK, ea=F_TITLE, lat=F_TITLE)
    p2 = para(tf)
    r2 = p2.add_run(); r2.text = "github.com/zai-org/ZCode"
    set_font(r2, 15, TEAL, ea=F_SANS, lat=F_SERIF)
    p2.space_before = Pt(12)
    line(s, M, Inches(3.30), Inches(5.6), RULE)
    p3 = para(tf)
    r3 = p3.add_run(); r3.text = "把代码交给社区监督，让 ZCode 变得开放、透明。"
    set_font(r3, 15, BODY, ea=F_SANS, lat=F_SANS)
    p3.space_before = Pt(22)
    p4 = para(tf)
    r4 = p4.add_run(); r4.text = "公告同时表示，欢迎开发者伙伴持续检查和反馈问题。"
    set_font(r4, 12.5, MUTED, ea=F_SANS, lat=F_SANS)
    p4.space_before = Pt(10)

    # ---------------------------------------------------------- 6 两份审计
    s = blank(prs)
    page_head(s, "04", "审计", "两份第三方安全审计")
    data = [
        ("审计机构", "中国信息通信研究院", "绿盟科技"),
        ("存储桶结论", "zcode-prod 阿里云 OSS 存储桶状态为云端零数据",
         "zcode-prod（阿里云 OSS）内全部数据对象及存储桶本身已删除"),
        ("客户端结论", "ZCode v3.14.0 客户端已完成安全整改；已移除 Repo Wiki，已切断本地仓库快照生成与上传链路",
         "ZCode v3.14.0 客户端已完成整改；Repo Wiki 入口及生成链路已移除，未发现可触发本地仓库快照或文件外发的功能路径"),
    ]
    cols = [Inches(2.05), Inches(4.75), Inches(4.78)]
    xx = M
    ytop = Inches(2.06)
    # 表头
    for j, (c, w) in enumerate(zip(cols, cols)):
        pass
    hdr = ["", "中国信息通信研究院", "绿盟科技"]
    xx = M
    for j, w in enumerate(cols):
        _, tf = tb(s, xx + Inches(0.16), ytop, w - Inches(0.2), Inches(0.34))
        p = para(tf, True)
        r = p.add_run(); r.text = hdr[j]
        set_font(r, 12.5, AMBER if j else MUTED, bold=bool(j), ea=F_SANS, lat=F_SANS)
        xx += w
    line(s, M, ytop + Inches(0.40), CW, INK, Pt(1.2))

    yy = ytop + Inches(0.56)
    for i, row in enumerate(data[1:], start=1):
        xx = M
        for j, w in enumerate(cols):
            txt = row[j]
            _, tf = tb(s, xx + Inches(0.16), yy, w - Inches(0.24), Inches(1.0))
            p = para(tf, True)
            r = p.add_run(); r.text = txt
            set_font(r, 11.5 if j else 12, BODY if j else MUTED,
                     bold=False, ea=F_SANS, lat=F_SANS)
            p.line_spacing = 1.3
            xx += w
        yy += Inches(1.30)
        line(s, M, yy - Inches(0.22), CW)

    figure_caption(s, M, Inches(6.62), CW,
                   "表 1 · 两份审计的结论并列（原文摘录）。审计覆盖对象见下页。")

    # ---------------------------------------------------------- 7 结论的边界
    s = blank(prs)
    page_head(s, "05", "解读", "审计证实了什么，公告承诺了什么")
    # 左：审计背书
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, M, Inches(2.10), Inches(5.5), Inches(3.5))
    box.fill.solid(); box.fill.fore_color.rgb = RGBColor.from_string("E8EEEC")
    box.line.color.rgb = RGBColor.from_string(TEAL); box.line.width = Pt(1)
    box.shadow.inherit = False
    _, tf = tb(s, M + Inches(0.36), Inches(2.36), Inches(4.8), Inches(3.0))
    p = para(tf, True)
    r = p.add_run(); r.text = "有第三方审计背书"
    set_font(r, 15, TEAL, bold=True, ea=F_SANS, lat=F_SANS)
    for t in ["zcode-prod 阿里云 OSS 存储桶：云端零数据",
              "存储桶内全部数据对象及存储桶本身已删除",
              "ZCode v3.14.0 客户端：已完成安全整改",
              "未发现可触发本地仓库快照或文件外发的功能路径"]:
        pp = para(tf)
        rr = pp.add_run(); rr.text = "·  " + t
        set_font(rr, 12.5, BODY, ea=F_SANS, lat=F_SANS)
        pp.space_before = Pt(9)
        pp.line_spacing = 1.3

    # 右：官方自述
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, M + Inches(6.0), Inches(2.10),
                             Inches(5.6), Inches(3.5))
    box.fill.solid(); box.fill.fore_color.rgb = RGBColor.from_string("F3EDE2")
    box.line.color.rgb = RGBColor.from_string(AMBER); box.line.width = Pt(1)
    box.shadow.inherit = False
    _, tf = tb(s, M + Inches(6.36), Inches(2.36), Inches(4.9), Inches(3.0))
    p = para(tf, True)
    r = p.add_run(); r.text = "出自公告的官方承诺"
    set_font(r, 15, AMBER, bold=True, ea=F_SANS, lat=F_SANS)
    for t in ["对于社区中提及的代码数据，我们承诺无留存",
              "也从未将其用于模型训练",
              "将建立常态化的产品安全漏洞机制",
              "按问题严重程度给予相应回报"]:
        pp = para(tf)
        rr = pp.add_run(); rr.text = "·  " + t
        set_font(rr, 12.5, BODY, ea=F_SANS, lat=F_SANS)
        pp.space_before = Pt(9)
        pp.line_spacing = 1.3

    _, tf = tb(s, M, Inches(5.86), CW, Inches(0.7))
    p = para(tf, True)
    r = p.add_run()
    r.text = ("读法：两份审计的结论文字只覆盖 OSS 存储桶与 v3.14.0 客户端两项；"
              "「无留存、从未用于模型训练」出自公告中的官方表述，审计结论未涉及。"
              "二者不是同一层面的证据。")
    set_font(r, 12.5, INK, ea=F_SANS, lat=F_SANS)
    p.line_spacing = 1.4

    # ---------------------------------------------------------- 8 承诺与机制
    s = blank(prs)
    page_head(s, "06", "承诺", "承诺与后续机制")
    fig = os.path.join(FIG, "04-mechanism.png")
    if os.path.exists(fig):
        s.shapes.add_picture(fig, M + Inches(6.55), Inches(2.02), width=Inches(5.05))
        figure_caption(s, M + Inches(6.55), Inches(4.95), Inches(5.05),
                       "图 2 · 常态化机制：发现—评估—回报—再检查的闭环")
    _, tf = tb(s, M, Inches(2.14), Inches(5.9), Inches(3.4))
    blocks = [
        ("数据承诺", "对于社区中提及的代码数据，承诺无留存，也从未将其用于模型训练。"),
        ("常态化机制", "建立常态化的产品安全漏洞机制，欢迎开发者伙伴持续检查和反馈问题。"),
        ("回报", "根据问题严重程度给予相应回报。"),
    ]
    first = True
    for t, d in blocks:
        p = para(tf, first); first = False
        r = p.add_run(); r.text = t
        set_font(r, 12, AMBER, bold=True, ea=F_SANS, lat=F_SANS, spacing=2.0)
        p2 = para(tf)
        r2 = p2.add_run(); r2.text = d
        set_font(r2, 14.5, INK, ea=F_SANS, lat=F_SANS)
        p2.space_before = Pt(6)
        p2.line_spacing = 1.35
        if t != "回报":
            p3 = para(tf)
            p3.space_before = Pt(20)

    # ---------------------------------------------------------- 9 流程线
    s = blank(prs)
    page_head(s, "07", "脉络", "从社区反馈到持续监督")
    steps = ["社区反馈问题", "完成整改并致歉", "开源代码", "第三方审计", "持续监督"]
    xx = M
    w = (CW - Inches(0.4) * 4) / 5
    for i, t in enumerate(steps):
        circ = s.shapes.add_shape(MSO_SHAPE.OVAL, xx + w / 2 - Inches(0.15),
                                  Inches(2.44), Inches(0.30), Inches(0.30))
        circ.fill.solid()
        circ.fill.fore_color.rgb = RGBColor.from_string(AMBER if i < 4 else TEAL)
        circ.line.fill.background(); circ.shadow.inherit = False
        _, tf = tb(s, xx, Inches(2.92), w, Inches(0.7))
        p = para(tf, True); p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = t
        set_font(r, 13, INK, bold=True, ea=F_SANS, lat=F_SANS)
        if i < 4:
            ar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, xx + w, Inches(2.585),
                                    Inches(0.4), Pt(1.2))
            ar.fill.solid(); ar.fill.fore_color.rgb = RGBColor.from_string(RULE)
            ar.line.fill.background(); ar.shadow.inherit = False
        xx += w + Inches(0.4)

    line(s, M, Inches(3.92), CW)
    _, tf = tb(s, M, Inches(4.16), CW, Inches(1.6))
    p = para(tf, True)
    r = p.add_run(); r.text = "公告中的对应表述"
    set_font(r, 12, MUTED, ea=F_SANS, lat=F_SANS, spacing=2.0)
    quotes = [
        "「针对社区反馈的 ZCode 产品安全问题，我们已经完成整改，并向所有用户道歉。」",
        "「我们已将 ZCode 开源，把代码交给社区监督，让 ZCode 变得开放、透明。」",
        "「经中国信息通信研究院技术评测……经绿盟科技审查……」",
        "「再次向大家致歉，请大家持续监督。」",
    ]
    for q in quotes:
        pp = para(tf)
        rr = pp.add_run(); rr.text = q
        set_font(rr, 12.5, BODY, ea=F_SANS, lat=F_SANS)
        pp.space_before = Pt(9)

    # ---------------------------------------------------------- 10 结语
    s = blank(prs, DARK)
    fig = os.path.join(FIG, "03-audit.png")
    if os.path.exists(fig):
        s.shapes.add_picture(fig, Inches(6.4), Inches(1.15), width=Inches(6.2))
        ov = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
        ov.fill.solid(); ov.fill.fore_color.rgb = RGBColor.from_string(DARK)
        ov.fill.transparency = 0.42
        ov.line.fill.background(); ov.shadow.inherit = False
    _, tf = tb(s, M, Inches(2.46), Inches(8.4), Inches(1.4))
    p = para(tf, True)
    r = p.add_run(); r.text = "再次向大家致歉"
    set_font(r, 40, WHITE, ea=F_TITLE, lat=F_TITLE, spacing=1.5)
    p2 = para(tf)
    r2 = p2.add_run(); r2.text = "请大家持续监督"
    set_font(r2, 40, "C8A15E", ea=F_TITLE, lat=F_TITLE, spacing=1.5)
    p2.space_before = Pt(2)

    line(s, M, Inches(4.90), Inches(1.5), AMBER, Pt(2.2))
    _, tf = tb(s, M, Inches(5.16), Inches(8.6), Inches(0.9))
    p = para(tf, True)
    r = p.add_run()
    r.text = "审计背书的是存储桶与客户端两项；「无留存、未用于训练」是公告的承诺。"
    set_font(r, 13, "A8B4C0", ea=F_SANS, lat=F_SANS)
    p2 = para(tf)
    r2 = p2.add_run(); r2.text = "来源：ZCode 官方公告（2026-09-21 08:45） ·  github.com/zai-org/ZCode"
    set_font(r2, 11.5, "7E8C9B", ea=F_SANS, lat=F_SERIF)
    p2.space_before = Pt(10)

    out = os.path.join(OUTDIR, "ZCode开源-官方公告解读.pptx")
    prs.save(out)
    print("已生成:", out)
    print("页数:", len(prs.slides.__iter__.__self__._sldIdLst))
    return out


if __name__ == "__main__":
    build()
