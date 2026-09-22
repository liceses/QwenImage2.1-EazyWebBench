#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
把 Qwen-Image-2.1 直接生成的整页幻灯片图片，打包成真正的 .pptx。

每页 = 一张满版图片（16:9 → 13.333x7.5 英寸），保持原始分辨率与版面。
用系统 Python 3.14 跑（它有 python-pptx）。
"""
import os
import sys
import glob

from pptx import Presentation
from pptx.util import Inches

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, ".."))
SRC = os.path.join(ROOT, "ppt_zcode", "gen_slides")
OUT = os.path.join(ROOT, "ppt_zcode", "ZCode开源-模型直出.pptx")

SW = Inches(13.333)
SH = Inches(7.5)


def main():
    files = sorted(glob.glob(os.path.join(SRC, "S*.png")))
    if not files:
        print("没有找到 S*.png 幻灯片图")
        return 1

    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    blank = prs.slide_layouts[6]

    for f in files:
        s = prs.slides.add_slide(blank)
        s.shapes.add_picture(f, 0, 0, width=SW, height=SH)
        print(f"  + {os.path.basename(f)}")

    prs.save(OUT)
    print(f"\n已生成: {OUT}")
    print(f"页数: {len(files)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
